"""Módulo de Registro y Diario de Operaciones en Excel.

Registra de forma limpia y transparente las señales de Compra, Venta,
Niveles de Stop Loss, Take Profit, Rendimientos calculados por el motor,
y el estado diario de "¿Requiere Acción Hoy?".
"""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd


class ExcelTradeJournal:
    """Gestiona un libro Excel profesional con el historial, niveles macro y órdenes pendientes."""

    def __init__(self, excel_path: str | Path | None = None):
        if excel_path is None:
            self.excel_path = Path(__file__).resolve().parents[2] / "data" / "diario_trading_oro.xlsx"
        else:
            self.excel_path = Path(excel_path)
        self.excel_path.parent.mkdir(parents=True, exist_ok=True)

    def export_operations(
        self,
        fills: list[dict[str, Any]],
        market_analysis: list[dict[str, Any]],
        portfolio_summary: dict[str, Any] | None = None,
        daily_status: list[dict[str, Any]] | None = None,
        control_zones: list[dict[str, Any]] | None = None,
    ) -> Path:
        """Exporta operaciones, análisis macro, estado del día y resumen a un archivo Excel multi-pestaña."""
        # 1. Pestaña de Estado Diario y Señales de Acción (si se proporciona)
        df_daily = pd.DataFrame(daily_status) if daily_status else pd.DataFrame()

        # 2. Pestaña de Operaciones y Señales
        rows_ops = []
        for fill in fills:
            rows_ops.append({
                "Fecha": fill.get("Fecha") or fill.get("date", datetime.now().strftime("%Y-%m-%d")),
                "Símbolo": fill.get("Símbolo") or fill.get("symbol"),
                "Tipo": str(fill.get("Tipo") or fill.get("side", "")).upper(),
                "Cantidad (Acciones)": int(fill.get("Cantidad (Acciones)") or fill.get("quantity", 0)),
                "Precio ($ USD)": float(fill.get("Precio ($ USD)") or fill.get("price", 0.0)),
                "Valor Total ($ USD)": float(fill.get("Valor Total ($ USD)") or round(float(fill.get("quantity", 0)) * float(fill.get("price", 0.0)), 2)),
                "Comisión ($ USD)": float(fill.get("Comisión ($ USD)") or fill.get("commission", 0.0)),
                "Stop Loss Sugerido ($)": float(fill.get("Stop Loss Sugerido ($)") or fill.get("stop_loss", 0.0)) if (fill.get("Stop Loss Sugerido ($)") or fill.get("stop_loss")) else "-",
                "Take Profit Sugerido ($)": float(fill.get("Take Profit Sugerido ($)") or fill.get("take_profit", 0.0)) if (fill.get("Take Profit Sugerido ($)") or fill.get("take_profit")) else "-",
                "P&L Operación ($ USD)": fill.get("P&L Operación ($ USD)", "-"),
                "Motivo / Señal": fill.get("Motivo / Señal") or fill.get("reason", "Señal Macro"),
            })

        df_ops = pd.DataFrame(rows_ops) if rows_ops else pd.DataFrame(columns=[
            "Fecha", "Símbolo", "Tipo", "Cantidad (Acciones)", "Precio ($ USD)",
            "Valor Total ($ USD)", "Comisión ($ USD)", "Stop Loss Sugerido ($)",
            "Take Profit Sugerido ($)", "P&L Operación ($ USD)", "Motivo / Señal"
        ])

        # 3. Pestaña de Análisis de Mercado (Soportes y Resistencias Reales)
        df_analysis = pd.DataFrame(market_analysis)

        # 4. Pestaña de Zonas de Control / Gaps Históricos Institucionales
        df_control_zones = pd.DataFrame(control_zones) if control_zones else pd.DataFrame()

        # 5. Pestaña de Resumen de Cartera
        resumen_list = []
        if portfolio_summary:
            for k, v in portfolio_summary.items():
                resumen_list.append({
                    "Métrica": k,
                    "Valor": v if isinstance(v, (int, float, Decimal)) else str(v),
                })
        df_summary = pd.DataFrame(resumen_list)

        target_file = self.excel_path
        attempts = [
            self.excel_path,
            self.excel_path.parent / f"{self.excel_path.stem}_nuevo.xlsx",
            self.excel_path.parent / f"{self.excel_path.stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
        ]
        written = False
        last_error = None
        for candidate in attempts:
            try:
                with pd.ExcelWriter(candidate, engine="openpyxl") as writer:
                    if not df_daily.empty:
                        df_daily.to_excel(writer, sheet_name="Estado Hoy y Señales", index=False)
                    df_ops.to_excel(writer, sheet_name="Operaciones y Registro", index=False)
                    df_analysis.to_excel(writer, sheet_name="Análisis de Niveles Macro", index=False)
                    if not df_control_zones.empty:
                        df_control_zones.to_excel(writer, sheet_name="Zonas de Control (Gaps)", index=False)
                    if not df_summary.empty:
                        df_summary.to_excel(writer, sheet_name="Resumen de Cartera", index=False)
                target_file = candidate
                written = True
                if candidate != self.excel_path:
                    print(f"\n[AVISO] {self.excel_path.name} está bloqueado por Excel.")
                    print(f"[AVISO] Guardando versión actualizada en: {candidate.name}")
                break
            except PermissionError as pe:
                last_error = pe
                continue

        if not written:
            raise PermissionError(f"No fue posible guardar el archivo Excel en {self.excel_path}: {last_error}")

        # Aplicar diseño y estilos openpyxl
        self._apply_openpyxl_styles(target_file)
        return target_file

    def _apply_openpyxl_styles(self, file_path: Path) -> None:
        """Aplica estilos visuales, colores, anchos y formatos numéricos en el libro Excel."""
        wb = openpyxl.load_workbook(file_path)

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        border_thin = Side(style="thin", color="E0E0E0")
        cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

        action_buy_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        action_buy_font = Font(name="Calibri", size=11, bold=True, color="276A3C")

        action_short_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        action_short_font = Font(name="Calibri", size=11, bold=True, color="C00000")

        action_hold_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
        action_hold_font = Font(name="Calibri", size=11, bold=True, color="1F4E79")

        action_wait_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        action_wait_font = Font(name="Calibri", size=11, bold=True, color="B25E00")

        pnl_pos_font = Font(name="Calibri", size=11, bold=True, color="276A3C")
        pnl_neg_font = Font(name="Calibri", size=11, bold=True, color="C00000")

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            if ws.max_row == 0:
                continue

            # Mostrar líneas de cuadrícula
            ws.views.sheetView[0].showGridLines = True
            ws.freeze_panes = "A2"

            # Formatear encabezado
            ws.row_dimensions[1].height = 28
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            # Mapear nombres de columnas
            headers = [str(ws.cell(row=1, column=c).value or "") for c in range(1, ws.max_column + 1)]

            is_summary_sheet = (ws.title == "Resumen de Cartera")

            # Formatear filas de datos
            for row_idx in range(2, ws.max_row + 1):
                ws.row_dimensions[row_idx].height = 20
                metric_name = str(ws.cell(row=row_idx, column=1).value or "") if is_summary_sheet else ""

                for col_idx in range(1, ws.max_column + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.border = cell_border
                    val = cell.value
                    header = headers[col_idx - 1]

                    if is_summary_sheet and col_idx == 1:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                        if any(k in str(val or "") for k in ["TOTAL", "VALOR TOTAL", "P&L NETO TOTAL"]):
                            cell.font = Font(name="Calibri", size=11, bold=True)
                        continue

                    # Determinar descriptor relevante (nombre de métrica en fila para Resumen, o encabezado de columna)
                    descriptor = metric_name if (is_summary_sheet and col_idx == 2) else header

                    # Formato numérico según columna o métrica
                    if isinstance(val, (int, float, Decimal)):
                        # 1. Comprobar porcentajes PRIMERO para evitar que "Distancia a Soporte (%)" sea tratada como "Soporte" ($)
                        if "%" in descriptor or "Rentabilidad" in descriptor or "Magnitud" in descriptor:
                            cell.value = float(val) / 100.0
                            cell.number_format = "+0.00%;-0.00%;0.00%" if any(k in descriptor for k in ["Var", "Rentabilidad"]) else "0.00%"
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                        # 2. Ratios Beneficio / Riesgo
                        elif any(k in descriptor for k in ["Ratio", "B/R", "R:R"]):
                            cell.value = float(val)
                            cell.number_format = '0.00":1"'
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                        # 3. Cantidad de acciones
                        elif any(k in descriptor for k in ["Cantidad", "Acciones", "Accs"]):
                            cell.number_format = "#,##0"
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                        # 4. Moneda y valores financieros en USD
                        elif any(k in descriptor for k in ["$", "Precio", "Valor", "Comisión", "Comisiones", "Soporte", "Resistencia", "SMA", "Stop Loss", "Take Profit", "P&L", "VALOR", "TOTAL", "BENEFICIO", "Límite"]):
                            cell.number_format = '"$"#,##0.00'
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                            # Colores y énfasis para P&L y totales
                            if "P&L" in descriptor or "BENEFICIO" in descriptor:
                                cell.font = pnl_pos_font if val >= 0 else pnl_neg_font
                            elif any(k in descriptor for k in ["TOTAL", "VALOR TOTAL", "P&L NETO"]):
                                cell.font = Font(name="Calibri", size=11, bold=True)
                        else:
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        # Texto / Fechas / Flags
                        if "Fecha" in header or "Símbolo" in header:
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                        elif any(k in header for k in ["Acción Hoy", "Dictamen", "Oportunidad", "Veredicto", "Señal", "Tipo", "Estado", "Rol"]):
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                            val_str = str(val or "").upper()
                            if any(k in val_str for k in ["TAKE PROFIT", "COMPRAR", "LONG", "ALCISTA", "SOPORTE"]):
                                cell.fill = action_buy_fill
                                cell.font = action_buy_font
                            elif any(k in val_str for k in ["STOP LOSS", "CORTO", "SHORT", "VENDER", "BAJISTA", "RESISTENCIA"]):
                                cell.fill = action_short_fill
                                cell.font = action_short_font
                            elif "MANTENER" in val_str or "ABIERTO" in val_str:
                                cell.fill = action_hold_fill
                                cell.font = action_hold_font
                            else:
                                cell.fill = action_wait_fill
                                cell.font = action_wait_font
                        else:
                            cell.alignment = Alignment(horizontal="left", vertical="center")

            # Ajustar anchos de columnas
            for col_idx in range(1, ws.max_column + 1):
                col_letter = get_column_letter(col_idx)
                header_text = headers[col_idx - 1]
                max_content_len = len(header_text)
                for r in range(2, ws.max_row + 1):
                    c_val = ws.cell(row=r, column=col_idx).value
                    if c_val is not None:
                        max_content_len = max(max_content_len, len(str(c_val)))

                ws.column_dimensions[col_letter].width = min(max(max_content_len + 4, 13), 60)

        wb.save(file_path)
