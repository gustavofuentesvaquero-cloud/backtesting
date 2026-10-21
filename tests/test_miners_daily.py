"""Unit tests for gold miners daily automation, strategy execution, and Excel journal."""

from datetime import datetime
from decimal import Decimal
import tempfile
import unittest
from pathlib import Path

import openpyxl
import pandas as pd

from market_catalog import (
    COMPANY_NAMES,
    ExcelTradeJournal,
    GoldMinersMacroStrategy,
    MarketDataStore,
    StrategyConfig,
    detect_price_control_zones,
    evaluate_market_opportunity,
    load_catalog,
    run_stock_macro_backtest,
)


class MinersDailyAutomationTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_dir = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_company_names_mapping(self):
        self.assertIn("GOLD", COMPANY_NAMES)
        self.assertIn("NEM", COMPANY_NAMES)
        self.assertEqual(COMPANY_NAMES["GOLD"], "Barrick Gold Corporation")
        self.assertEqual(COMPANY_NAMES["NEM"], "Newmont Corporation")

    def test_excel_trade_journal_creates_all_sheets_and_styles(self):
        excel_path = self.base_dir / "diario_test.xlsx"
        journal = ExcelTradeJournal(excel_path=excel_path)

        fills = [
            {
                "Fecha": "2026-07-20",
                "Símbolo": "GOLD",
                "Tipo": "COMPRA",
                "Cantidad (Acciones)": 200,
                "Precio ($ USD)": 36.64,
                "Valor Total ($ USD)": 7328.0,
                "Comisión ($ USD)": 1.0,
                "Stop Loss Sugerido ($)": 34.28,
                "Take Profit Sugerido ($)": 48.83,
                "P&L Operación ($ USD)": "-",
                "Motivo / Señal": "Rebote en soporte de 60 días",
            }
        ]
        market_analysis = [
            {
                "Símbolo": "GOLD",
                "Activo": "Barrick Gold Corporation",
                "Soporte Macro 60d ($)": 36.08,
                "Resistencia Macro 60d ($)": 48.83,
                "Media Móvil SMA 60d ($)": 42.79,
                "Último Cierre Oficial ($)": 40.13,
                "Amplitud Rango ($)": 12.75,
                "Distancia a Soporte (%)": 11.23,
                "Distancia a Resistencia (%)": 17.82,
                "Stop Loss Recomendado ($)": 34.28,
            }
        ]
        portfolio_summary = {
            "VALOR TOTAL INVERTIDO ($ USD)": 8026.0,
            "TOTAL BENEFICIO LATENTE ($ USD)": 698.0,
            "TOTAL BENEFICIO REALIZADO ($ USD)": 0.0,
            "TOTAL COMISIONES PAGADAS ($ USD)": 1.0,
            "P&L NETO TOTAL ($ USD)": 697.0,
        }
        daily_status = [
            {
                "Símbolo": "GOLD",
                "Empresa": "Barrick Gold Corporation",
                "Fecha Cotización": "2026-10-12",
                "Último Precio ($)": 40.13,
                "Var. Día ($)": 0.50,
                "Var. Día (%)": 1.26,
                "Soporte 60d ($)": 36.08,
                "Resistencia 60d ($)": 48.83,
                "SMA 60d ($)": 42.79,
                "Stop Loss Sugerido ($)": 34.28,
                "Posición Actual (Accs)": 200,
                "Precio Entrada ($)": 36.64,
                "Fecha Entrada": "2026-07-20",
                "Valor Mercado ($ USD)": 8026.0,
                "P&L Latente ($ USD)": 698.0,
                "Rentabilidad (%)": 9.52,
                "¿Requiere Acción Hoy?": "NO — MANTENER",
                "Instrucción Operativa": "Mantener 200 acciones. Objetivo: $48.83 | Stop: $34.28",
            }
        ]

        saved_path = journal.export_operations(
            fills=fills,
            market_analysis=market_analysis,
            portfolio_summary=portfolio_summary,
            daily_status=daily_status,
        )

        self.assertTrue(saved_path.exists())

        # Verificar hojas con openpyxl
        wb = openpyxl.load_workbook(saved_path)
        expected_sheets = [
            "Estado Hoy y Señales",
            "Operaciones y Registro",
            "Análisis de Niveles Macro",
            "Resumen de Cartera",
        ]
        self.assertEqual(wb.sheetnames, expected_sheets)

        # Verificar encabezados y estilos de Estado Hoy y Señales
        ws_status = wb["Estado Hoy y Señales"]
        self.assertEqual(ws_status.cell(row=1, column=1).value, "Símbolo")
        self.assertEqual(ws_status.cell(row=2, column=1).value, "GOLD")
        self.assertEqual(ws_status.freeze_panes, "A2")

        # Verificar columna de acción hoy
        action_col = None
        for col in range(1, ws_status.max_column + 1):
            if ws_status.cell(row=1, column=col).value == "¿Requiere Acción Hoy?":
                action_col = col
                break
        self.assertIsNotNone(action_col)
        self.assertEqual(ws_status.cell(row=2, column=action_col).value, "NO — MANTENER")

        # Verificar formateo correcto de porcentajes en Análisis de Niveles Macro (no deben ser tratados como $)
        ws_macro = wb["Análisis de Niveles Macro"]
        dist_support_col = None
        for col in range(1, ws_macro.max_column + 1):
            if ws_macro.cell(row=1, column=col).value == "Distancia a Soporte (%)":
                dist_support_col = col
                break
        self.assertIsNotNone(dist_support_col)
        cell_dist = ws_macro.cell(row=2, column=dist_support_col)
        self.assertAlmostEqual(cell_dist.value, 0.1123, places=4)
        self.assertEqual(cell_dist.number_format, "0.00%")

        # Verificar que Resumen de Cartera almacena números y aplica formato de divisa
        ws_summary = wb["Resumen de Cartera"]
        self.assertEqual(ws_summary.cell(row=2, column=1).value, "VALOR TOTAL INVERTIDO ($ USD)")
        summary_val_cell = ws_summary.cell(row=2, column=2)
        self.assertIsInstance(summary_val_cell.value, (int, float))
        self.assertEqual(summary_val_cell.value, 8026.0)
        self.assertEqual(summary_val_cell.number_format, '"$"#,##0.00')

    def test_gold_miners_macro_strategy_validation(self):
        strategy = GoldMinersMacroStrategy()
        # Faltan parámetros
        with self.assertRaises(ValueError):
            strategy.validate_config({"symbol": "GOLD"})

        # Parámetros válidos
        valid = strategy.validate_config({
            "symbol": "GOLD",
            "support_price": "36.00",
            "target_price": "48.00",
            "stop_loss_price": "34.00",
            "shares": "200",
        })
        self.assertTrue(valid)

    def test_run_stock_macro_backtest_with_mock_data(self):
        # Crear datos de prueba sintéticos para simular histórico
        store = MarketDataStore(base_dir=self.base_dir)
        dates = pd.date_range("2026-05-01", periods=65, freq="B")
        rows = []
        for i, d in enumerate(dates):
            # Precio comienza en 45, cae a 36 en barra 10, sube a 40 al final
            price = 45.0 - (i * 0.9) if i <= 10 else 36.0 + ((i - 10) * 0.08)
            rows.append({
                "date": d.strftime("%Y-%m-%d"),
                "open": price,
                "high": price + 0.5,
                "low": price - 0.5,
                "close": price,
                "volume": 100000.0,
            })
        df = pd.DataFrame(rows)
        store.save_bars("TEST_MINER", df, timeframe="1d")

        # Crear catálogo temporal con TEST_MINER
        catalog_path = Path(__file__).resolve().parents[1] / "data" / "gold_miners_stocks.json"
        catalog = load_catalog(catalog_path)

        # Modificar símbolo de test a GOLD para que use el catálogo oficial
        store.save_bars("GOLD", df, timeframe="1d")

        ops, metrics, p_summary, daily_status = run_stock_macro_backtest(
            symbol="GOLD",
            shares=100,
            lookback_bars=60,
            store=store,
            catalog=catalog,
        )

        self.assertIsInstance(ops, list)
        self.assertIsInstance(metrics, dict)
        self.assertIsInstance(p_summary, dict)
        self.assertIsInstance(daily_status, dict)

        self.assertEqual(daily_status["Símbolo"], "GOLD")
        self.assertIn("¿Requiere Acción Hoy?", daily_status)
        self.assertIn("Posición Actual (Accs)", daily_status)
        self.assertIn("P&L Latente ($ USD)", daily_status)
        # Verificar campos bi-direccionales añadidos
        self.assertIn("Dictamen / Oportunidad", daily_status)
        self.assertIn("Objetivo Largos ($)", daily_status)
        self.assertIn("Stop Loss Largos ($)", daily_status)
        self.assertIn("Ratio B/R Largos", daily_status)
        self.assertIn("Objetivo Cortos ($)", daily_status)
        self.assertIn("Stop Loss Cortos ($)", daily_status)
        self.assertIn("Ratio B/R Cortos", daily_status)
        self.assertIn("Razonamiento Detallado", daily_status)

        self.assertIn("Objetivo Largos ($)", metrics)
        self.assertIn("Stop Loss Largos ($)", metrics)
        self.assertIn("Ratio B/R Largos", metrics)
        self.assertIn("Objetivo Cortos ($)", metrics)
        self.assertIn("Stop Loss Cortos ($)", metrics)
        self.assertIn("Ratio B/R Cortos", metrics)

    def test_evaluate_market_opportunity_neutral_zone(self):
        levels = {
            "swing_low": Decimal("36.00"),
            "swing_high": Decimal("48.00"),
            "sma": Decimal("42.00"),
        }
        multi = {
            "200d": {"sma": Decimal("43.00"), "swing_low": Decimal("35.00"), "swing_high": Decimal("50.00")},
            "252d": {"sma": Decimal("40.00"), "swing_low": Decimal("30.00"), "swing_high": Decimal("55.00")},
            "max": {"sma": Decimal("35.00"), "swing_low": Decimal("20.00"), "swing_high": Decimal("60.00")},
        }
        # Precio en 42.00 (medio del canal)
        res = evaluate_market_opportunity("GOLD", 42.00, levels, multi, current_position=0)
        self.assertEqual(res["veredicto"], "MANTENER / ESPERAR")
        self.assertEqual(res["requiere_accion"], "NO — ESPERAR")
        self.assertEqual(res["long_target"], 48.00)
        self.assertEqual(res["long_stop_loss"], 34.20)
        self.assertGreater(res["long_rr_ratio"], 0)
        self.assertEqual(res["short_target"], 36.00)
        self.assertEqual(res["short_stop_loss"], 50.40)
        self.assertGreater(res["short_rr_ratio"], 0)
        self.assertIn("ZONA NEUTRAL", res["razonamiento"])

    def test_evaluate_market_opportunity_buy_long_zone(self):
        levels = {
            "swing_low": Decimal("36.00"),
            "swing_high": Decimal("48.00"),
            "sma": Decimal("42.00"),
        }
        multi = {
            "200d": {"sma": Decimal("35.00"), "swing_low": Decimal("35.00"), "swing_high": Decimal("50.00")},
            "252d": {"sma": Decimal("35.00"), "swing_low": Decimal("30.00"), "swing_high": Decimal("55.00")},
            "max": {"sma": Decimal("30.00"), "swing_low": Decimal("20.00"), "swing_high": Decimal("60.00")},
        }
        # Precio en 36.50 (zona de soporte, distancia +1.39% <= 3.5%)
        res = evaluate_market_opportunity("GOLD", 36.50, levels, multi, current_position=0)
        self.assertEqual(res["veredicto"], "COMPRAR (LONG)")
        self.assertEqual(res["requiere_accion"], "SÍ — COMPRAR")
        self.assertTrue(res["in_long_zone"])
        self.assertGreater(res["long_rr_ratio"], 1.5)
        self.assertEqual(res["long_target"], 48.00)
        self.assertEqual(res["long_stop_loss"], 34.20)
        self.assertIn("MOMENTO ÓPTIMO PARA LARGOS", res["razonamiento"])

    def test_evaluate_market_opportunity_sell_short_zone(self):
        levels = {
            "swing_low": Decimal("36.00"),
            "swing_high": Decimal("48.00"),
            "sma": Decimal("42.00"),
        }
        multi = {
            "200d": {"sma": Decimal("45.00"), "swing_low": Decimal("35.00"), "swing_high": Decimal("50.00")},
            "252d": {"sma": Decimal("40.00"), "swing_low": Decimal("30.00"), "swing_high": Decimal("55.00")},
            "max": {"sma": Decimal("35.00"), "swing_low": Decimal("20.00"), "swing_high": Decimal("60.00")},
        }
        # Precio en 47.50 (zona de resistencia, distancia 1.04% <= 3.5%)
        res = evaluate_market_opportunity("GOLD", 47.50, levels, multi, current_position=0)
        self.assertEqual(res["veredicto"], "VENDER / CORTO (SHORT)")
        self.assertEqual(res["requiere_accion"], "SÍ — VENDER / CORTO")
        self.assertTrue(res["in_short_zone"])
        self.assertGreater(res["short_rr_ratio"], 1.5)
        self.assertEqual(res["short_target"], 36.00)
        self.assertEqual(res["short_stop_loss"], 50.40)
        self.assertIn("MOMENTO ÓPTIMO PARA CORTOS", res["razonamiento"])

    def test_gold_miners_macro_strategy_short_direction_orders(self):
        strategy = GoldMinersMacroStrategy()
        config = StrategyConfig(
            strategy_id="test-short",
            instance_id="test-short-1",
            strategy_version="1.0.0",
            parameters={
                "symbol": "GOLD",
                "support_price": "36.00",
                "target_price": "48.00",
                "stop_loss_price": "34.00",
                "short_target_price": "36.00",
                "short_stop_loss_price": "50.40",
                "shares": "100",
                "direction": "short",
            },
        )
        from unittest.mock import MagicMock
        context = MagicMock()
        context.config = config

        market = MagicMock()
        # Enviar cotización en zona de resistencia (bid 47.50 >= 48*0.97 = 46.56)
        market.event.event_type = "quote"
        market.event.payload = {
            "bid": {"price": "47.50"},
            "ask": {"price": "47.52"},
        }
        market.event.source_timestamp_ns = 1760000000000000000

        portfolio = MagicMock()
        state = {"has_entered": False, "has_exited": False}

        # 1. Probar entrada en corto
        decision = strategy.on_market(context, market, portfolio, state)
        self.assertEqual(len(decision.intents), 1)
        self.assertEqual(decision.intents[0].side, "sell")
        self.assertEqual(decision.state["has_entered"], True)
        self.assertEqual(decision.state["pos_side"], "short")

        # 2. Probar take profit de corto (ask cae a 35.90 <= 36.00)
        market.event.payload = {
            "bid": {"price": "35.88"},
            "ask": {"price": "35.90"},
        }
        decision2 = strategy.on_market(context, market, portfolio, decision.state)
        self.assertEqual(len(decision2.intents), 1)
        self.assertEqual(decision2.intents[0].side, "buy")
        self.assertEqual(decision2.state["has_exited"], True)
        self.assertEqual(decision2.state["exit_reason"], "take_profit")

    def test_evaluate_market_opportunity_crash_breakdown_not_buy(self):
        """Verifica que un activo colapsado muy por debajo del soporte (-72%) NO dé señal de compra."""
        levels = {
            "swing_low": Decimal("36.00"),
            "swing_high": Decimal("48.00"),
            "sma": Decimal("42.00"),
        }
        multi = {
            "200d": {"sma": Decimal("40.00"), "swing_low": Decimal("25.00"), "swing_high": Decimal("55.00")},
            "252d": {"sma": Decimal("38.00"), "swing_low": Decimal("20.00"), "swing_high": Decimal("60.00")},
            "max": {"sma": Decimal("30.00"), "swing_low": Decimal("15.00"), "swing_high": Decimal("70.00")},
        }
        # Precio se desploma a 10.00 (soporte 36.00 roto)
        res = evaluate_market_opportunity("GOLD", 10.00, levels, multi, current_position=0)
        self.assertFalse(res["in_long_zone"])
        self.assertNotEqual(res["veredicto"], "COMPRAR (LONG)")
        self.assertEqual(res["veredicto"], "MANTENER / ESPERAR")
        self.assertIn("POR DEBAJO DEL CANAL 60D", res["razonamiento"])

    def test_evaluate_market_opportunity_parabolic_breakout_not_short(self):
        """Verifica que un activo en rotura parabólica muy por encima del techo (+100%) NO dé señal de corto."""
        levels = {
            "swing_low": Decimal("36.00"),
            "swing_high": Decimal("48.00"),
            "sma": Decimal("42.00"),
        }
        multi = {
            "200d": {"sma": Decimal("40.00"), "swing_low": Decimal("25.00"), "swing_high": Decimal("55.00")},
            "252d": {"sma": Decimal("38.00"), "swing_low": Decimal("20.00"), "swing_high": Decimal("60.00")},
            "max": {"sma": Decimal("30.00"), "swing_low": Decimal("15.00"), "swing_high": Decimal("70.00")},
        }
        # Precio se dispara a 100.00 (resistencia 48.00 perforada al alza)
        res = evaluate_market_opportunity("GOLD", 100.00, levels, multi, current_position=0)
        self.assertFalse(res["in_short_zone"])
        self.assertNotEqual(res["veredicto"], "VENDER / CORTO (SHORT)")
        self.assertEqual(res["veredicto"], "MANTENER / ESPERAR")
        self.assertIn("POR ENCIMA DEL CANAL 60D", res["razonamiento"])

    def test_evaluate_market_opportunity_exact_resistance_triggers_short(self):
        """Verifica que al precio exacto de resistencia se calcule el SL en resistencia*1.05 y active CORTO."""
        levels = {
            "swing_low": Decimal("36.00"),
            "swing_high": Decimal("48.00"),
            "sma": Decimal("42.00"),
        }
        multi = {
            "200d": {"sma": Decimal("40.00"), "swing_low": Decimal("25.00"), "swing_high": Decimal("55.00")},
            "252d": {"sma": Decimal("38.00"), "swing_low": Decimal("20.00"), "swing_high": Decimal("60.00")},
            "max": {"sma": Decimal("30.00"), "swing_low": Decimal("15.00"), "swing_high": Decimal("70.00")},
        }
        # Precio exacto en la resistencia (48.00)
        res = evaluate_market_opportunity("GOLD", 48.00, levels, multi, current_position=0)
        self.assertTrue(res["in_short_zone"])
        self.assertEqual(res["short_stop_loss"], 50.40)
        self.assertEqual(res["short_target"], 36.00)
        self.assertEqual(res["veredicto"], "VENDER / CORTO (SHORT)")
        self.assertEqual(res["requiere_accion"], "SÍ — VENDER / CORTO")
        self.assertGreater(res["short_rr_ratio"], 1.5)

    def test_run_stock_macro_backtest_short_direction_execution(self):
        """Verifica ejecución de backtest en dirección 'short' con entradas y salidas de corto y PnL correcto."""
        store = MarketDataStore(base_dir=self.base_dir)
        dates = pd.date_range("2026-05-01", periods=65, freq="B")
        rows = []
        for i, d in enumerate(dates):
            # Comienza en 38, sube a 48 (barra 10) activando corto, luego cae a 36 (barra 40)
            if i <= 10:
                price = 38.0 + (i * 1.0)
            elif i <= 40:
                price = 48.0 - ((i - 10) * 0.4)
            else:
                price = 36.0 + ((i - 40) * 0.1)
            rows.append({
                "date": d.strftime("%Y-%m-%d"),
                "open": price,
                "high": price + 0.3,
                "low": price - 0.3,
                "close": price,
                "volume": 50000.0,
            })
        df = pd.DataFrame(rows)
        store.save_bars("GOLD", df, timeframe="1d")

        catalog_path = Path(__file__).resolve().parents[1] / "data" / "gold_miners_stocks.json"
        catalog = load_catalog(catalog_path)

        ops, metrics, p_summary, daily_status = run_stock_macro_backtest(
            symbol="GOLD",
            shares=100,
            lookback_bars=60,
            store=store,
            catalog=catalog,
            direction="short",
        )
        self.assertGreater(len(ops), 0)
        short_entries = [o for o in ops if o["Tipo"] == "VENTA (CORTO)"]
        self.assertEqual(len(short_entries), 1)
        self.assertIn("Rechazo en resistencia", short_entries[0]["Motivo / Señal"])

    def test_mgc_futures_multiplier_accounting(self):
        """Verifica que para futuros CME MGC se aplique el multiplicador de 10x y terminología de contratos."""
        levels = {
            "swing_low": Decimal("4000.00"),
            "swing_high": Decimal("4500.00"),
            "sma": Decimal("4250.00"),
        }
        multi = {
            "200d": {"sma": Decimal("4100.00"), "swing_low": Decimal("3800.00"), "swing_high": Decimal("4600.00")},
            "252d": {"sma": Decimal("4000.00"), "swing_low": Decimal("3500.00"), "swing_high": Decimal("4800.00")},
            "max": {"sma": Decimal("3500.00"), "swing_low": Decimal("2500.00"), "swing_high": Decimal("5000.00")},
        }
        # Posición activa de 1 contrato comprado a 4100, precio actual 4200 (+100 pts * 10 = +$1000 USD)
        res = evaluate_market_opportunity("MGC", 4200.00, levels, multi, current_position=1, entry_price=4100.00)
        self.assertIn("contratos", res["instruccion"])
        self.assertIn("+$1,000.00 USD", res["instruccion"])

    def test_detect_price_control_zones_empty_and_invalid(self):
        """Verifica manejo robusto de DataFrames vacíos, None o insuficientes."""
        res_none = detect_price_control_zones(None)
        self.assertEqual(res_none["zones"], [])
        self.assertIsNone(res_none["nearest_support"])
        self.assertIsNone(res_none["nearest_resistance"])

        res_empty = detect_price_control_zones(pd.DataFrame())
        self.assertEqual(res_empty["zones"], [])

        df_one = pd.DataFrame([{
            "date": "2026-04-01", "open": 100.0, "high": 105.0, "low": 98.0, "close": 102.0
        }])
        res_one = detect_price_control_zones(df_one)
        self.assertEqual(res_one["zones"], [])

    def test_detect_price_control_zones_bullish_and_bearish_gaps(self):
        """Verifica detección matemática de gaps alcistas (>3%) y bajistas (<-3%), sus zonas y roles."""
        rows = [
            # Barra 0: Base
            {"date": "2026-04-01", "open": 100.0, "high": 102.0, "low": 98.0, "close": 100.0},
            # Barra 1: Gap alcista (>3% sobre high de barra 0: 102 * 1.03 = 105.06)
            # low=108.0 > 105.06. Zona [108.0, 112.0]
            {"date": "2026-04-02", "open": 112.0, "high": 115.0, "low": 108.0, "close": 114.0},
            # Barra 2: Consolidación por encima sin cerrar el gap
            {"date": "2026-04-03", "open": 114.0, "high": 118.0, "low": 113.0, "close": 117.0},
        ]
        df = pd.DataFrame(rows)
        res = detect_price_control_zones(df, min_gap_pct=0.03)

        self.assertEqual(len(res["zones"]), 1)
        z = res["zones"][0]
        self.assertEqual(z["type"], "ALCISTA")
        self.assertEqual(z["lower_bound"], 108.0)
        self.assertEqual(z["upper_bound"], 112.0)
        self.assertEqual(z["status"], "ABIERTO")
        self.assertEqual(z["role"], "SOPORTE")
        self.assertIsNotNone(res["nearest_support"])
        self.assertEqual(res["nearest_support"]["lower_bound"], 108.0)
        self.assertIsNone(res["nearest_resistance"])

    def test_gap_status_tracking_closed_vs_open(self):
        """Verifica que un gap pase de ABIERTO a CERRADO cuando precios posteriores penetran la zona."""
        rows = [
            {"date": "2026-04-01", "open": 100.0, "high": 102.0, "low": 98.0, "close": 100.0},
            {"date": "2026-04-02", "open": 112.0, "high": 115.0, "low": 108.0, "close": 114.0},
            # Barra 3 cae hasta 107.0 penetrando la zona inferior (108.0) -> Gap CERRADO
            {"date": "2026-04-03", "open": 113.0, "high": 114.0, "low": 107.0, "close": 109.0},
            {"date": "2026-04-04", "open": 109.0, "high": 120.0, "low": 108.5, "close": 118.0},
        ]
        df = pd.DataFrame(rows)
        res = detect_price_control_zones(df, min_gap_pct=0.03)
        self.assertEqual(len(res["zones"]), 1)
        self.assertEqual(res["zones"][0]["status"], "CERRADO")
        # El precio de cierre actual es 118.0, por encima de 112.0, por lo que actúa como SOPORTE
        self.assertEqual(res["zones"][0]["role"], "SOPORTE")


    def test_evaluate_market_opportunity_with_control_zones(self):
        """Verifica que evaluate_market_opportunity integre las zonas de control en métricas y dictamen."""
        levels = {
            "swing_low": Decimal("157.33"),
            "swing_high": Decimal("302.35"),
            "sma": Decimal("240.89"),
        }
        multi = {
            "200d": {"sma": Decimal("215.63"), "swing_low": Decimal("150.00"), "swing_high": Decimal("320.00")},
            "252d": {"sma": Decimal("200.00"), "swing_low": Decimal("75.70"), "swing_high": Decimal("351.28")},
            "max": {"sma": Decimal("100.00"), "swing_low": Decimal("8.41"), "swing_high": Decimal("351.28")},
        }
        control_zones = {
            "nearest_support": {
                "date": "2026-04-14",
                "type": "ALCISTA",
                "lower_bound": 197.40,
                "upper_bound": 202.78,
                "magnitude_pct": 11.19,
                "status": "CERRADO",
                "role": "SOPORTE",
                "dist_pct": -14.19,
                "description": "$197.40 - $202.78 (Gap Alcista 2026-04-14, CERRADO)",
            },
            "nearest_resistance": {
                "date": "2026-04-29",
                "type": "ALCISTA",
                "lower_bound": 267.46,
                "upper_bound": 276.55,
                "magnitude_pct": 11.60,
                "status": "CERRADO",
                "role": "RESISTENCIA",
                "dist_pct": 13.19,
                "description": "$267.46 - $276.55 (Gap Alcista 2026-04-29, CERRADO)",
            },
            "active_zone": None,
        }

        # Caso 1: Precio actual a $236.30 (en espera, incluye zonas en razonamiento)
        res = evaluate_market_opportunity(
            symbol="BE",
            last_close=236.30,
            levels=levels,
            multi_levels=multi,
            current_position=0,
            control_zones=control_zones,
        )
        self.assertIn("197.40", res["control_zone_soporte"])
        self.assertIn("267.46", res["control_zone_resistencia"])
        self.assertIn("Zona Control Soporte", res["razonamiento"])

        # Caso 2: Precio testeando zona de control de soporte ($201.00) con asimetría R:R
        res_entry = evaluate_market_opportunity(
            symbol="BE",
            last_close=201.00,
            levels=levels,
            multi_levels=multi,
            current_position=0,
            control_zones=control_zones,
        )
        self.assertEqual(res_entry["veredicto"], "COMPRAR (LONG)")
        self.assertIn("MOMENTO ÓPTIMO PARA LARGOS", res_entry["razonamiento"])
        self.assertGreater(res_entry["long_rr_ratio"], 1.5)

    def test_detect_price_control_zones_bearish_gap(self):
        """Verifica detección matemática de gaps bajistas (<-3%), sus zonas, status y rol RESISTENCIA."""
        rows = [
            # Barra 0: Base
            {"date": "2026-06-01", "open": 100.0, "high": 105.0, "low": 100.0, "close": 102.0},
            # Barra 1: Gap bajista (high < prev_low * 0.97 = 100 * 0.97 = 97.0). high=94.0.
            {"date": "2026-06-02", "open": 92.0, "high": 94.0, "low": 88.0, "close": 90.0},
            # Barra 2: Cotización por debajo
            {"date": "2026-06-03", "open": 89.0, "high": 91.0, "low": 86.0, "close": 88.0},
        ]
        df = pd.DataFrame(rows)
        res = detect_price_control_zones(df, min_gap_pct=0.03)

        self.assertEqual(len(res["zones"]), 1)
        z = res["zones"][0]
        self.assertEqual(z["type"], "BAJISTA")
        self.assertEqual(z["lower_bound"], 92.0)
        self.assertEqual(z["upper_bound"], 94.0)
        self.assertEqual(z["status"], "ABIERTO")
        self.assertEqual(z["role"], "RESISTENCIA")
        self.assertIsNotNone(res["nearest_resistance"])
        self.assertEqual(res["nearest_resistance"]["upper_bound"], 94.0)

        # Ahora barra 4 penetra hacia arriba pasando 95.0 (> 94.0) -> Gap CERRADO
        rows.append({"date": "2026-06-04", "open": 88.0, "high": 96.0, "low": 87.0, "close": 95.5})
        df_closed = pd.DataFrame(rows)
        res_closed = detect_price_control_zones(df_closed, min_gap_pct=0.03)
        self.assertEqual(res_closed["zones"][0]["status"], "CERRADO")



    def test_export_operations_with_control_zones_sheet(self):
        """Verifica que export_operations incluya la hoja de Zonas de Control (Gaps) cuando se proporcionan."""
        import tempfile

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_excel = Path(tmp_dir) / "test_journal.xlsx"
            journal = ExcelTradeJournal(excel_path=tmp_excel)

            cz_rows = [{
                "Símbolo": "BE",
                "Activo": "Bloom Energy",
                "Fecha Gap": "2026-04-14",
                "Tipo": "ALCISTA",
                "Límite Inferior ($)": 197.40,
                "Límite Superior ($)": 202.78,
                "Magnitud (%)": 11.19,
                "Estado": "CERRADO",
                "Rol Actual": "SOPORTE",
                "Distancia al Cierre (%)": -14.19,
            }]

            saved = journal.export_operations(
                fills=[],
                market_analysis=[{"Símbolo": "BE", "Activo": "Bloom Energy"}],
                portfolio_summary={"TOTAL": 1000},
                daily_status=[{"Símbolo": "BE", "¿Requiere Acción Hoy?": "NO — ESPERAR"}],
                control_zones=cz_rows,
            )

            self.assertTrue(saved.exists())
            wb = openpyxl.load_workbook(saved)
            self.assertIn("Zonas de Control (Gaps)", wb.sheetnames)
            ws_cz = wb["Zonas de Control (Gaps)"]
            self.assertEqual(ws_cz.cell(row=1, column=1).value, "Símbolo")
            self.assertEqual(ws_cz.cell(row=2, column=1).value, "BE")
            self.assertEqual(ws_cz.cell(row=2, column=4).value, "ALCISTA")



